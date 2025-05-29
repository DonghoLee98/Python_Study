from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
# B2:C11 까지의 데이터를 chart로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()          # chart 종류 설정 (Bar, Line, Pie, ...)
# bar_chart.add_data(bar_value)   # chart에 데이터 추가
# ws.add_chart(bar_chart, "E1")

# B1:C11 까지의 데이터를 chart로 생성
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True)  # titles_from_data 를 True로 하면 Legend 제목을 정의할 수 있다
line_chart.title = "성적표" # 제목
line_chart.style = 10   # 미리 정의된 스타일 적용. (사용자가 개별 지정도 가능)
line_chart.y_axis.title = "점수"    # y축의 제목
line_chart.x_axis.title = "번호"    # x축의 제목
ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")

# openpyxl 모듈 Library
# https://openpyxl.readthedocs.io/en/stable/api/openpyxl.chart.html#module-openpyxl.chart
# 위 페이지에서 라이브러리 확인 후 사용하면 될 것